"""Google Play Reviews Scraper
==============================

Collect user reviews from the Google Play Store and save them as CSV, JSON
or Excel. Works as a library (``GooglePlayReviewScraper``) and as a CLI::

    python src/review_scraper.py com.whatsapp
    python src/review_scraper.py "https://play.google.com/store/apps/details?id=com.whatsapp" \
        --country us --lang en --max 5000 --since 2026-01-01 --format csv,xlsx

Run it with no arguments for an interactive prompt.
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import sys
import threading
import time
from collections.abc import Iterable
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Optional
from urllib.parse import parse_qs, urlparse

from google_play_scraper import Sort, app, reviews

__all__ = [
    "GooglePlayReviewScraper",
    "Review",
    "ScrapeResult",
    "SUPPORTED_FORMATS",
    "parse_app_id",
    "is_valid_app_id",
    "clean_filename",
]

log = logging.getLogger("gplay_reviews")

SUPPORTED_FORMATS = ("csv", "json", "xlsx")

_APP_ID_RE = re.compile(r"^[a-zA-Z][a-zA-Z0-9_]*(\.[a-zA-Z][a-zA-Z0-9_]*)+$")

# (collected_so_far, expected_total_or_None)
ProgressCallback = Callable[[int, Optional[int]], None]


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def parse_app_id(text: str) -> str:
    """Return the app id from a Play Store URL or a bare package name.

    Raises ``ValueError`` if nothing that looks like an app id is found.
    """
    text = (text or "").strip()
    if not text:
        raise ValueError("empty input")

    if "://" in text or text.startswith("play.google.com"):
        url = text if "://" in text else f"https://{text}"
        app_id = parse_qs(urlparse(url).query).get("id", [None])[0]
        if not app_id:
            match = re.search(r"[?&]id=([^&#\s]+)", text)
            app_id = match.group(1) if match else None
        if not app_id:
            raise ValueError(f"no app id found in URL: {text}")
        text = app_id

    if not is_valid_app_id(text):
        raise ValueError(f"invalid app id: {text}")
    return text


def is_valid_app_id(app_id: str) -> bool:
    """Check that ``app_id`` looks like an Android package name."""
    return bool(_APP_ID_RE.match(app_id or ""))


def clean_filename(name: str, max_len: int = 40) -> str:
    """Make ``name`` safe to use as a file name on every OS."""
    cleaned = re.sub(r"\s+", "_", name or "")
    cleaned = re.sub(r"[^\w\-.]", "", cleaned)
    cleaned = re.sub(r"_+", "_", cleaned).strip("._")
    return cleaned[:max_len].rstrip("._").lower() or "app"


def _iso(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return "" if value is None else str(value)


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #
@dataclass
class Review:
    """Simplified representation of a Play Store review."""

    review_id: str
    user_name: str
    content: str
    score: int
    thumbs_up_count: int
    at: str
    reply_content: str
    reply_at: str
    app_version: str

    @classmethod
    def from_raw(cls, raw: dict) -> Review:
        """Build a Review from the dict returned by ``google_play_scraper``."""
        return cls(
            review_id=raw.get("reviewId") or "",
            user_name=raw.get("userName") or "",
            content=raw.get("content") or "",
            score=int(raw.get("score") or 0),
            thumbs_up_count=int(raw.get("thumbsUpCount") or 0),
            at=_iso(raw.get("at")),
            reply_content=raw.get("replyContent") or "",
            reply_at=_iso(raw.get("repliedAt")),
            app_version=raw.get("appVersion") or raw.get("reviewCreatedVersion") or "",
        )

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class ScrapeResult:
    """Outcome of collecting one app."""

    app_id: str
    app_info: dict
    reviews: list[Review] = field(default_factory=list)
    complete: bool = True          # False if stopped early by error or cancel
    cancelled: bool = False
    error: Optional[str] = None

    @property
    def title(self) -> str:
        return self.app_info.get("title") or self.app_id

    @property
    def expected_total(self) -> Optional[int]:
        total = self.app_info.get("reviews")
        return int(total) if isinstance(total, (int, float)) and total > 0 else None


# --------------------------------------------------------------------------- #
# Scraper
# --------------------------------------------------------------------------- #
class GooglePlayReviewScraper:
    """High-level interface for extracting reviews from Google Play."""

    def __init__(
        self,
        country: str = "br",
        lang: str = "pt",
        *,
        batch_size: int = 1000,
        delay: float = 0.3,
        max_retries: int = 4,
        empty_page_retries: int = 2,
    ) -> None:
        self.country = country.lower()
        self.lang = lang.lower()
        self.batch_size = batch_size
        self.delay = delay
        self.max_retries = max_retries
        self.empty_page_retries = empty_page_retries

    # -- kept for backwards compatibility -------------------------------- #
    def extract_app_id_from_url(self, url: str) -> str:
        return parse_app_id(url)

    def validate_app_id(self, app_id: str) -> bool:
        return is_valid_app_id(app_id)

    def clean_filename(self, filename: str) -> str:
        return clean_filename(filename)

    # -------------------------------------------------------------------- #
    def get_app_info(self, app_id: str) -> dict:
        """Return app metadata; never raises (falls back to a stub)."""
        try:
            info = self._with_retries(
                lambda: app(app_id, lang=self.lang, country=self.country), None
            )
            return info or {"appId": app_id, "title": app_id}
        except Exception as exc:  # noqa: BLE001 - metadata is best-effort
            log.warning("Could not fetch app info for %s: %s", app_id, exc)
            return {"appId": app_id, "title": app_id, "error": str(exc)}

    def scrape(
        self,
        app_id: str,
        *,
        max_reviews: Optional[int] = None,
        since: Optional[date] = None,
        cancel_event: Optional[threading.Event] = None,
        on_progress: Optional[ProgressCallback] = None,
    ) -> ScrapeResult:
        """Collect reviews for ``app_id`` (newest first).

        - ``max_reviews``: stop after this many reviews.
        - ``since``: stop at the first review older than this date.
        - ``cancel_event``: set it from another thread to stop between batches.
        - ``on_progress(collected, expected_total)``: called after every batch.

        Network errors are retried with exponential backoff. If retries run
        out, the reviews collected so far are kept and ``complete`` is False.
        """
        info = self.get_app_info(app_id)
        result = ScrapeResult(app_id=app_id, app_info=info)

        expected = result.expected_total
        if max_reviews:
            expected = min(expected, max_reviews) if expected else max_reviews
        since_dt = datetime.combine(since, datetime.min.time()) if since else None

        token = None
        batch = 0
        log.info("Collecting %s (%s)", result.title, app_id)

        while True:
            if cancel_event is not None and cancel_event.is_set():
                result.complete, result.cancelled = False, True
                log.info("Cancelled after %d reviews", len(result.reviews))
                break

            batch += 1
            want = self.batch_size
            if max_reviews:
                want = min(want, max_reviews - len(result.reviews))

            try:
                raw, token = self._fetch_batch(app_id, want, token, cancel_event)
            except _Cancelled:
                result.complete, result.cancelled = False, True
                break
            except Exception as exc:  # noqa: BLE001
                result.complete, result.error = False, str(exc)
                log.error("Batch %d failed after %d retries: %s", batch, self.max_retries, exc)
                break

            stop = False
            for item in raw:
                if since_dt is not None:
                    at = item.get("at")
                    if isinstance(at, datetime) and at < since_dt:
                        stop = True
                        break
                result.reviews.append(Review.from_raw(item))
                if max_reviews and len(result.reviews) >= max_reviews:
                    stop = True
                    break

            log.debug("Batch %d: %d reviews (total %d)", batch, len(raw), len(result.reviews))
            if on_progress:
                on_progress(len(result.reviews), expected)

            if stop or not raw or token is None or token.token is None:
                break
            if cancel_event is not None:
                cancel_event.wait(self.delay)
            else:
                time.sleep(self.delay)

        if not result.reviews and not result.error and info.get("error"):
            # Metadata failed too: most likely a network problem or a bad app id.
            result.complete, result.error = False, info["error"]

        log.info("Done: %d reviews for %s", len(result.reviews), app_id)
        return result

    def extract_all_reviews(self, app_id: str, **kwargs) -> tuple[list[Review], dict]:
        """Backwards-compatible wrapper around :meth:`scrape`."""
        result = self.scrape(app_id, **kwargs)
        return result.reviews, result.app_info

    def _fetch_batch(self, app_id: str, count: int, token, cancel_event):
        """Fetch one page of reviews.

        ``google_play_scraper.reviews`` swallows network errors and returns an
        empty page with no continuation token - the same signal as "no more
        reviews". So an empty page is retried with the *same* token a few
        times before we accept it as the real end.
        """
        def call():
            return reviews(
                app_id,
                lang=self.lang,
                country=self.country,
                sort=Sort.NEWEST,
                count=count,
                continuation_token=token,
            )

        for attempt in range(self.empty_page_retries + 1):
            raw, next_token = self._with_retries(call, cancel_event)
            if raw:
                return raw, next_token
            if attempt < self.empty_page_retries:
                wait = 2 ** attempt
                log.debug("Empty page, retrying in %ss to rule out a network hiccup", wait)
                self._sleep(wait, cancel_event)
        return [], None

    @staticmethod
    def _sleep(seconds: float, cancel_event: Optional[threading.Event]) -> None:
        if cancel_event is None:
            time.sleep(seconds)
        elif cancel_event.wait(seconds):
            raise _Cancelled()

    def _with_retries(self, fn, cancel_event: Optional[threading.Event]):
        for attempt in range(self.max_retries + 1):
            try:
                return fn()
            except Exception as exc:  # noqa: BLE001
                if attempt >= self.max_retries:
                    raise
                wait = min(2 ** attempt, 30)
                log.warning("Request failed (%s); retrying in %ss", exc, wait)
                self._sleep(wait, cancel_event)
        raise RuntimeError("unreachable")

    # -------------------------------------------------------------------- #
    def save_reviews(
        self,
        reviews_data: Iterable[Review | dict],
        app_info: Optional[dict] = None,
        filename_prefix: Optional[str] = None,
        *,
        output_dir: str | Path = ".",
        formats: Iterable[str] = ("csv", "json"),
    ) -> dict[str, Path]:
        """Save reviews and return ``{format: path}`` for each file written."""
        rows = [r.to_dict() if isinstance(r, Review) else dict(r) for r in reviews_data]
        if not rows:
            log.info("No reviews to save")
            return {}

        formats = [f.lower().strip() for f in formats]
        unknown = set(formats) - set(SUPPORTED_FORMATS)
        if unknown:
            raise ValueError(f"unsupported format(s): {', '.join(sorted(unknown))}")

        app_info = app_info or {}
        if not filename_prefix:
            title = app_info.get("title") or app_info.get("appId")
            filename_prefix = f"{clean_filename(title)}_reviews" if title else "google_play_reviews"

        out = Path(output_dir).expanduser()
        out.mkdir(parents=True, exist_ok=True)
        base = out / filename_prefix
        written: dict[str, Path] = {}

        columns = list(dict.fromkeys(key for row in rows for key in row))

        if "csv" in formats:
            path = base.with_suffix(".csv")
            # utf-8-sig so Excel opens accents correctly
            with open(path, "w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=columns, restval="")
                writer.writeheader()
                writer.writerows(rows)
            written["csv"] = path

        if "xlsx" in formats:
            path = base.with_suffix(".xlsx")
            _write_xlsx(path, columns, rows)
            written["xlsx"] = path

        if "json" in formats:
            path = base.with_suffix(".json")
            payload = {
                "collection_info": {
                    "collected_at": datetime.now().isoformat(timespec="seconds"),
                    "total_reviews": len(rows),
                    "country": self.country,
                    "language": self.lang,
                },
                "app_info": app_info,
                "reviews": rows,
            }
            with open(path, "w", encoding="utf-8") as fh:
                json.dump(payload, fh, ensure_ascii=False, indent=2, default=str)
            written["json"] = path

        for path in written.values():
            log.info("Saved %s", path)
        return written

    # -------------------------------------------------------------------- #
    @staticmethod
    def statistics(reviews_data: Iterable[Review | dict]) -> dict:
        """Return count, average score and the 1-5 star distribution."""
        scores = [
            (r.score if isinstance(r, Review) else int(r.get("score") or 0))
            for r in reviews_data
        ]
        scores = [s for s in scores if 1 <= s <= 5]
        dist = {star: scores.count(star) for star in range(1, 6)}
        avg = sum(scores) / len(scores) if scores else 0.0
        return {"total": len(scores), "average": avg, "distribution": dist}


class _Cancelled(Exception):
    pass


# Characters Excel refuses inside a cell (control chars other than \t \n \r).
_XLSX_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_XLSX_WIDTHS = {"content": 80, "reply_content": 60, "user_name": 24, "review_id": 24,
                "at": 20, "reply_at": 20, "app_version": 14}


def _write_xlsx(path: Path, columns: list[str], rows: list[dict]) -> None:
    """Stream rows into an .xlsx with openpyxl's write-only mode (low memory)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("reviews")
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _XLSX_WIDTHS.get(col, 12)

    bold = Font(bold=True)
    header = []
    for col in columns:
        cell = WriteOnlyCell(ws, value=col)
        cell.font = bold
        header.append(cell)
    ws.append(header)

    for row in rows:
        values = []
        for col in columns:
            value = row.get(col, "")
            if isinstance(value, str):
                value = _XLSX_ILLEGAL.sub("", value)
            elif not isinstance(value, (int, float, type(None))):
                value = str(value)
            values.append(value)
        ws.append(values)
    wb.save(path)


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def _print_summary(result: ScrapeResult, files: dict[str, Path]) -> None:
    stats = GooglePlayReviewScraper.statistics(result.reviews)
    print(f"\n{result.title}: {stats['total']} reviews, average {stats['average']:.2f}★")
    for star in range(5, 0, -1):
        count = stats["distribution"][star]
        pct = count / stats["total"] * 100 if stats["total"] else 0
        bar = "█" * round(pct / 4)
        print(f"  {star}★ {bar:<25} {count:>7} ({pct:4.1f}%)")
    if not result.complete:
        reason = "cancelled" if result.cancelled else f"error: {result.error}"
        print(f"  ⚠ partial collection ({reason})")
    for path in files.values():
        print(f"  → {path}")


def _progress_printer(app_id: str) -> ProgressCallback:
    def show(done: int, total: Optional[int]) -> None:
        suffix = f" / ~{total:,}" if total else ""
        sys.stderr.write(f"\r  {app_id}: {done:,}{suffix} reviews")
        sys.stderr.flush()

    return show


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("use YYYY-MM-DD") from exc


def _parse_formats(value: str) -> list[str]:
    formats = [f.strip().lower() for f in value.split(",") if f.strip()]
    bad = [f for f in formats if f not in SUPPORTED_FORMATS]
    if bad or not formats:
        raise argparse.ArgumentTypeError(f"choose from {', '.join(SUPPORTED_FORMATS)}")
    return formats


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="review_scraper",
        description="Collect Google Play Store reviews into CSV, JSON or Excel.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""examples:
  review_scraper.py com.whatsapp
  review_scraper.py "https://play.google.com/store/apps/details?id=com.whatsapp"
  review_scraper.py com.whatsapp com.instagram.android --country us --lang en
  review_scraper.py com.whatsapp --max 1000 --since 2026-01-01 --format csv,xlsx""",
    )
    parser.add_argument("apps", nargs="*", metavar="APP", help="Play Store URL or app id (one or more)")
    parser.add_argument("--url", "-u", action="append", default=[], help=argparse.SUPPRESS)
    parser.add_argument("--app-id", "-a", action="append", default=[], help=argparse.SUPPRESS)
    parser.add_argument("--country", "-c", default="br", help="store country code (default: br)")
    parser.add_argument("--lang", "-l", default="pt", help="review language code (default: pt)")
    parser.add_argument("--max", "-n", type=int, dest="max_reviews", help="max reviews per app")
    parser.add_argument("--since", type=_parse_date, help="only reviews on/after YYYY-MM-DD")
    parser.add_argument("--format", "-f", type=_parse_formats, default=["csv", "json"],
                        help="comma-separated: csv,json,xlsx (default: csv,json)")
    parser.add_argument("--output-dir", "-d", default=".", help="folder for output files")
    parser.add_argument("--output", "-o", help="file name prefix (single app only)")
    parser.add_argument("--verbose", "-v", action="store_true", help="show debug logs")
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.WARNING,
        format="%(levelname)s: %(message)s",
    )

    inputs = [*args.apps, *args.url, *args.app_id]
    if not inputs:
        return interactive_mode()

    app_ids = []
    for item in inputs:
        try:
            app_ids.append(parse_app_id(item))
        except ValueError as exc:
            print(f"✗ {exc}", file=sys.stderr)
            return 2
    if args.output and len(app_ids) > 1:
        print("✗ --output only works with a single app", file=sys.stderr)
        return 2

    scraper = GooglePlayReviewScraper(country=args.country, lang=args.lang)
    exit_code = 0
    for app_id in app_ids:
        try:
            result = scraper.scrape(
                app_id,
                max_reviews=args.max_reviews,
                since=args.since,
                on_progress=_progress_printer(app_id),
            )
        except KeyboardInterrupt:
            print("\n✗ interrupted", file=sys.stderr)
            return 130
        sys.stderr.write("\n")
        if not result.reviews:
            print(f"✗ {app_id}: no reviews found ({result.error or 'empty'})", file=sys.stderr)
            exit_code = 1
            continue
        files = scraper.save_reviews(
            result.reviews, result.app_info, args.output,
            output_dir=args.output_dir, formats=args.format,
        )
        _print_summary(result, files)
    return exit_code


def interactive_mode() -> int:
    print("Google Play Reviews Scraper")
    print("─" * 40)
    try:
        raw = input("App URL or id: ").strip()
        app_id = parse_app_id(raw)
        country = input("Country [br]: ").strip() or "br"
        lang = input("Language [pt]: ").strip() or "pt"
        max_raw = input("Max reviews [all]: ").strip()
    except ValueError as exc:
        print(f"✗ {exc}")
        return 2
    except (KeyboardInterrupt, EOFError):
        print()
        return 130

    argv = [app_id, "--country", country, "--lang", lang]
    if max_raw.isdigit():
        argv += ["--max", max_raw]
    return main(argv)


if __name__ == "__main__":
    sys.exit(main())
